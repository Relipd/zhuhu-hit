# -*- coding: utf-8 -*-
"""回归验证:在临时 ROOT 上重跑交付链, 与既有交付物逐字节/逐单元格比对。

为什么必须有一支脚本(SKILL.md「回归方式」的外化)
------------------------------------------------
改了任何脚本后, 都要证明"产出没变"(而不是凭感觉)。手工做这件事要复制数据、跑五条命令、
逐个文件比字节、再比 Excel 单元格 —— 容易漏、容易忘。本脚本把它固化成一条命令:

    python regress_pipeline.py --root <ROOT> --date 2026-09-14

判据:
  · 临时 ROOT 上 `check → fill_excel → gen_html → verify_html → topic_lib rebuild` 全部 exit 0;
  · 根入口 HTML / 各详情页 / `话题库.md` **逐字节一致**;
  · 月度 Excel **当日 sheet 逐单元格一致**(2026-09-21 起; 临时根只留当日数据,
    跨日累积的 sheet 不参与比对 —— 整簿比对在多日月份下必然误报)。
**差异来自数据本身**(如既有交付物早于某片段最后一次写入, 见坑 69)不算失败 —— 此时先重跑
`run_pipeline.py` 刷新既有交付物, 再跑本脚本。

退出码: 0=通过; 1=有差异或某一关失败(逐条打印)。
"""
import argparse
import io
import os
import shutil
import subprocess
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import contract   # noqa: E402  数据契约:交付物命名的单一定义处

for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

COPY_DIRS = ("raw", "ext_search", contract.LIB_DIRNAME, contract.STYLE_DIR)


def run_chain(py, scripts, root, date):
    """在指定 ROOT 上跑交付链; 返回 (是否全通过, 各关退出码)。"""
    steps = [("check", ["check.py", "--root", root, "--date", date]),
             ("fill_excel", ["fill_excel.py", "--root", root, "--date", date]),
             ("gen_html", ["gen_html.py", "--root", root, "--date", date]),
             ("verify_html", ["verify_html.py", "--root", root, "--date", date]),
             ("topic_lib rebuild", ["topic_lib.py", "rebuild", "--root", root])]
    ok = True
    codes = {}
    for label, argv in steps:
        r = subprocess.run([py, os.path.join(scripts, argv[0])] + argv[1:],
                           capture_output=True, text=True, encoding="utf-8", errors="replace")
        codes[label] = r.returncode
        tail = ((r.stdout or "").strip().splitlines() or [""])[-1][:100]
        print("   %-18s exit=%d | %s" % (label, r.returncode, tail))
        if r.returncode != 0:
            ok = False
            print((r.stdout or "")[-800:])
            print((r.stderr or "")[-800:])
    return ok, codes


def compare_files(real_root, tmp_root, date):
    """逐字节比对交付物; 返回 (一致数, 差异列表, 缺失列表)。"""
    pairs = [("根入口", contract.report_entry(real_root, date), contract.report_entry(tmp_root, date)),
             ("话题库.md", os.path.join(real_root, contract.LIB_DIRNAME, contract.LIB_MD),
              os.path.join(tmp_root, contract.LIB_DIRNAME, contract.LIB_MD))]
    a_dir, b_dir = contract.report_pages_dir(real_root, date), contract.report_pages_dir(tmp_root, date)
    if os.path.isdir(a_dir) and os.path.isdir(b_dir):
        for name in sorted(os.listdir(a_dir)):
            pairs.append((name, os.path.join(a_dir, name), os.path.join(b_dir, name)))

    same, diff, missing = 0, [], []
    for label, a, b in pairs:
        if not os.path.exists(b):
            missing.append(label)
            continue
        if not os.path.exists(a):
            diff.append("%s(既有文件不存在)" % label)
            continue
        if io.open(a, "rb").read() == io.open(b, "rb").read():
            same += 1
        else:
            diff.append(label)
    return same, diff, missing


def compare_excel(real_root, tmp_root, date):
    """逐单元格比对月度 Excel 的**当日 sheet**(2026-09-21 起只比当日)。

    临时 ROOT 只保留当日数据(下方 --tmp 流程的减噪设计), 而月度工作簿的
    「发帖短评/热点拓展」sheet 与其他日期 sheet 跨日累积 —— 整簿比对必然误报。
    因此只对 sheet 名 == str(date) 的当日页做全量比对; 其余 sheet 明确跳过并注记。
    """
    try:
        from openpyxl import load_workbook
    except Exception as e:
        return -1, ["openpyxl 不可用: %s" % e]
    a_path, b_path = contract.xlsx_path(real_root, date), contract.xlsx_path(tmp_root, date)
    if not os.path.exists(b_path):
        return -1, ["临时 ROOT 未生成 %s" % os.path.basename(b_path)]
    if not os.path.exists(a_path):
        return -1, ["既有 %s 不存在" % os.path.basename(a_path)]
    wa, wb = load_workbook(a_path), load_workbook(b_path)
    today = str(date)
    if today not in wb.sheetnames:
        return -1, ["临时 ROOT 未生成当日 sheet: %s" % today]
    if today not in wa.sheetnames:
        return -1, ["既有工作簿没有当日 sheet: %s" % today]
    skipped = len(wa.sheetnames) - 1
    notes = ["跳过非当日 sheet %d 个(跨日累积, 单日临时根不可比)" % skipped] if skipped else []
    n = 0
    A, B = wa[today], wb[today]
    if (A.max_row, A.max_column) != (B.max_row, B.max_column):
        notes.append("%s 尺寸不同 %sx%s vs %sx%s"
                     % (today, A.max_row, A.max_column, B.max_row, B.max_column))
        return 1, notes
    for r in range(1, A.max_row + 1):
        for c in range(1, A.max_column + 1):
            if A.cell(r, c).value != B.cell(r, c).value:
                n += 1
                if len(notes) < 5:
                    notes.append("%s R%dC%d 不一致" % (today, r, c))
    return n, notes


def main():
    ap = argparse.ArgumentParser(description="交付链回归验证(临时 ROOT 重跑 + 逐字节/逐单元格比对)")
    ap.add_argument("--root", required=True)
    ap.add_argument("--date", required=True)
    ap.add_argument("--python", default=sys.executable, help="跑子进程用的解释器(默认当前解释器)")
    ap.add_argument("--tmp", default=None, help="临时 ROOT 位置(默认系统临时目录/track_regress)")
    ap.add_argument("--keep", action="store_true", help="保留临时 ROOT(默认跑完删除)")
    args = ap.parse_args()

    script_dir = os.path.dirname(os.path.abspath(__file__))
    tmp = args.tmp or os.path.join(tempfile.gettempdir(), "track_regress")
    if os.path.isdir(tmp):
        shutil.rmtree(tmp, ignore_errors=True)
    os.makedirs(tmp, exist_ok=True)
    for item in COPY_DIRS:
        src = os.path.join(args.root, item)
        if os.path.isdir(src):
            shutil.copytree(src, os.path.join(tmp, item))
    # 只保留当日数据, 减少拷贝与噪声
    for sub in ("raw", "ext_search"):
        d = os.path.join(tmp, sub)
        if os.path.isdir(d):
            for name in os.listdir(d):
                if name != args.date:
                    p = os.path.join(d, name)
                    shutil.rmtree(p, ignore_errors=True) if os.path.isdir(p) else os.remove(p)

    print("== 临时 ROOT 重跑交付链: %s" % tmp)
    chain_ok, _ = run_chain(args.python, script_dir, tmp, args.date)

    print("\n== 逐字节比对")
    same, diff, missing = compare_files(args.root, tmp, args.date)
    print("   文件: 一致 %d | 不一致 %s | 缺失 %s" % (same, diff or 0, missing or 0))

    print("\n== Excel 单元格比对")
    cell_diff, notes = compare_excel(args.root, tmp, args.date)
    for x in notes:
        print("   %s" % x)
    print("   单元格差异合计: %s" % cell_diff)

    ok = chain_ok and not diff and not missing and cell_diff == 0
    print("\n== 结论: %s" % ("通过(行为零变化)" if ok else "需排查(先看差异是否来自数据本身, 见坑 69)"))
    if not args.keep:
        shutil.rmtree(tmp, ignore_errors=True)
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
